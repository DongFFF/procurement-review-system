"""
采购需求文件解析。

支持格式：.docx / .pdf / .txt / .md
产出统一的「文档结构」：分节 + 段落，保留原始定位信息，
以便审查意见能精确回指到原文，修订时能精确替换。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any


class ParseError(RuntimeError):
    """文件解析失败。"""


# ---------------------------------------------------------------------------
# 数据结构
# ---------------------------------------------------------------------------

@dataclass
class Block:
    """文档中的一个内容块（段落 / 表格行 / 标题）。"""

    index: int                      # 在文档中的顺序号
    text: str                       # 文本内容
    kind: str = "paragraph"         # paragraph | heading | table | list
    level: int = 0                  # 标题层级
    section: str = ""               # 所属章节名（如「技术参数」）
    locator: dict[str, Any] = field(default_factory=dict)  # 定位信息

    def to_dict(self) -> dict[str, Any]:
        return {
            "index": self.index,
            "text": self.text,
            "kind": self.kind,
            "level": self.level,
            "section": self.section,
            "locator": self.locator,
        }


@dataclass
class ParsedDocument:
    """解析后的文档。"""

    filename: str
    filetype: str
    blocks: list[Block] = field(default_factory=list)
    meta: dict[str, Any] = field(default_factory=dict)

    @property
    def full_text(self) -> str:
        return "\n".join(b.text for b in self.blocks if b.text.strip())

    def sections(self) -> dict[str, str]:
        """按章节聚合文本，便于规则按部位检索。"""
        out: dict[str, list[str]] = {}
        for b in self.blocks:
            key = b.section or "正文"
            out.setdefault(key, []).append(b.text)
        return {k: "\n".join(v) for k, v in out.items()}

    def to_dict(self) -> dict[str, Any]:
        return {
            "filename": self.filename,
            "filetype": self.filetype,
            "meta": self.meta,
            "char_count": len(self.full_text),
            "block_count": len(self.blocks),
            "sections": list(self.sections().keys()),
            "blocks": [b.to_dict() for b in self.blocks],
        }


# ---------------------------------------------------------------------------
# 章节识别
# ---------------------------------------------------------------------------

# 采购需求文件常见章节关键词 → 标准部位名
SECTION_KEYWORDS: list[tuple[str, str]] = [
    (r"项目(概况|概述|背景|简介)", "项目概况"),
    (r"采购(需求|内容|标的|清单|范围)", "采购需求"),
    (r"(技术|功能)(参数|要求|规格|需求|指标)", "技术参数"),
    (r"(商务|合同)(要求|条款|条件)", "商务要求"),
    (r"资格(要求|条件)|供应商(资格|条件)|投标人资格", "资格条件"),
    (r"(评分|评审|评标)(办法|标准|细则|因素)|打分", "评分办法"),
    (r"(预算|最高限价|资金)(金额|情况|来源)?", "预算与资金"),
    (r"(交付|交货|工期|服务)(时间|期限|地点|要求)", "交付要求"),
    (r"(售后|质保|维保)(服务|要求|期)?", "售后服务"),
    (r"(验收|履约)(要求|标准|方案|办法)", "履约验收"),
    (r"(付款|支付)(方式|条件|进度|比例)", "付款方式"),
    (r"(培训|实施)(要求|方案|计划)?", "培训与实施"),
    (r"(中小|小微)企业", "中小企业政策"),
    (r"(节能|环保|绿色)(产品|采购)?", "节能环保"),
    (r"进口产品", "进口产品"),
]

_HEADING_PATTERNS = [
    # 第X章 / 第X节 / 第X部分 —— 几乎必然是标题
    (re.compile(r"^\s*第[一二三四五六七八九十百]+[章节部分]\s*[、.．]?\s*(.*)$"), True),
    # 一、二、三、 —— 通常是标题
    (re.compile(r"^\s*[一二三四五六七八九十]+\s*[、.．]\s*(.+)$"), True),
    # （一）（二） —— 通常是子标题
    (re.compile(r"^\s*[（(]\s*[一二三四五六七八九十]+\s*[)）]\s*(.+)$"), True),
    # 1. / 1.1 / 1、 —— 可能是标题，也可能是编号列表项，需进一步判断
    (re.compile(r"^\s*(\d+(?:\.\d+){0,3})\s*[、.．]?\s*(\S.{0,60})$"), False),
]

# 行尾出现这些标点，说明是完整句子而非标题
_SENTENCE_END = ("。", "；", "！", "？", ".", ";", "!", "?")


def _looks_like_list_item(t: str) -> bool:
    """
    判断一个带编号的短行到底是「标题」还是「正文列表项」。

    采购需求文件里 "1. CPU：Intel Xeon...，2 颗" 这类是正文条目，
    而 "四、评分办法" 这类才是标题。区分依据：
      - 以句号/分号结尾 → 正文
      - 长度较长 → 正文
      - 冒号后跟着大段描述 → 正文
    """
    if t.endswith(_SENTENCE_END):
        return True
    if len(t) > 25:
        return True
    # "内存：512GB DDR5 ECC RDIMM，频率不低于 4800MHz" 这种冒号后内容很长的
    if "：" in t or ":" in t:
        head, _, tail = t.replace(":", "：").partition("：")
        if len(tail) > 12:
            return True
    return False


def _detect_section(text: str) -> str | None:
    """
    判断一段文本是否是新章节标题，返回标准部位名或 None。

    设计原则：**只按标准部位切分**。
    规则库的 applies_to 用的是「技术参数」「资格条件」这类标准部位名，
    如果按 "（一）计算节点" 这类子标题切分，会产生大量碎片章节，
    导致规则匹配不到、RAG 上下文被割裂。

    因此：
      - 命中标准部位关键词 → 切分为对应部位
      - 未命中关键词的子标题（如"（一）计算节点"）→ 返回 None，归入上级部位
    """
    t = text.strip()
    if not t or len(t) > 60:
        return None

    # 最高优先级：命中标准部位关键词
    for kw, name in SECTION_KEYWORDS:
        if re.search(kw, t):
            # 排除正文句子误判：如"本项目预算金额为 860 万元。"含"预算"但非标题
            if _looks_like_list_item(t) and not re.match(
                r"^\s*(第[一二三四五六七八九十百]+[章节部分]|[一二三四五六七八九十]+\s*[、.．]|[（(]\s*[一二三四五六七八九十]+\s*[)）])",
                t,
            ):
                continue
            return name

    # 无编号的短行，仅当命中关键词（上面已处理）
    return None


# ---------------------------------------------------------------------------
# 各格式解析器
# ---------------------------------------------------------------------------

def _parse_docx(path: Path) -> tuple[list[Block], dict[str, Any]]:
    try:
        from docx import Document  # type: ignore
    except ImportError as e:
        raise ParseError("缺少 python-docx，无法解析 Word 文件") from e

    doc = Document(str(path))
    blocks: list[Block] = []
    idx = 0
    current_section = "正文"

    # 段落与标题，保持文档顺序
    from docx.oxml.ns import qn  # type: ignore
    from docx.table import Table  # type: ignore
    from docx.text.paragraph import Paragraph  # type: ignore

    body = doc.element.body
    para_map = {p._p: p for p in doc.paragraphs}
    table_map = {t._tbl: t for t in doc.tables}

    for child in body.iterchildren():
        if child in para_map:
            p = para_map[child]
            text = p.text.strip()
            if not text:
                continue
            style_name = (p.style.name or "") if p.style is not None else ""
            is_heading = style_name.lower().startswith("heading") or style_name.startswith("标题")
            level = 0
            if is_heading:
                m = re.search(r"(\d+)", style_name)
                level = int(m.group(1)) if m else 1

            sec = _detect_section(text)
            if sec:
                current_section = sec
                level = level or 1

            blocks.append(
                Block(
                    index=idx,
                    text=text,
                    kind="heading" if (is_heading or sec) else "paragraph",
                    level=level,
                    section=current_section,
                    locator={"type": "docx_paragraph", "para_index": len(blocks)},
                )
            )
            idx += 1

        elif child in table_map:
            t = table_map[child]
            for r_i, row in enumerate(t.rows):
                cells = [c.text.strip() for c in row.cells]
                # 去掉合并单元格造成的重复
                dedup: list[str] = []
                for c in cells:
                    if not dedup or dedup[-1] != c:
                        dedup.append(c)
                line = " | ".join(x for x in dedup if x)
                if not line:
                    continue
                blocks.append(
                    Block(
                        index=idx,
                        text=line,
                        kind="table",
                        section=current_section,
                        locator={"type": "docx_table", "row": r_i},
                    )
                )
                idx += 1

    meta = {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "has_tables": len(doc.tables) > 0,
    }
    return blocks, meta


def _parse_pdf(path: Path) -> tuple[list[Block], dict[str, Any]]:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError as e:
        raise ParseError("缺少 pypdf，无法解析 PDF 文件") from e

    reader = PdfReader(str(path))
    blocks: list[Block] = []
    idx = 0
    current_section = "正文"

    for page_no, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        for line in text.splitlines():
            line = line.strip()
            if not line:
                continue
            sec = _detect_section(line)
            if sec:
                current_section = sec
            blocks.append(
                Block(
                    index=idx,
                    text=line,
                    kind="heading" if sec else "paragraph",
                    level=1 if sec else 0,
                    section=current_section,
                    locator={"type": "pdf_page", "page": page_no + 1},
                )
            )
            idx += 1

    meta = {"pages": len(reader.pages)}
    if not blocks:
        meta["warning"] = (
            "未能从 PDF 中提取到文本，该文件可能是扫描件（图片型 PDF）。"
            "建议先用 OCR 转换，或改传 Word 版本。"
        )
    return blocks, meta


def _parse_text(path: Path) -> tuple[list[Block], dict[str, Any]]:
    raw = ""
    for enc in ("utf-8", "utf-8-sig", "gbk", "gb18030", "latin-1"):
        try:
            raw = path.read_text(encoding=enc)
            break
        except (UnicodeDecodeError, LookupError):
            continue
    if not raw:
        raise ParseError("无法识别文件编码，请另存为 UTF-8 后重试")

    blocks: list[Block] = []
    current_section = "正文"
    for idx, line in enumerate(raw.splitlines()):
        line = line.strip()
        if not line:
            continue
        sec = _detect_section(line)
        if sec:
            current_section = sec
        blocks.append(
            Block(
                index=idx,
                text=line,
                kind="heading" if sec else "paragraph",
                level=1 if sec else 0,
                section=current_section,
                locator={"type": "line", "line": idx + 1},
            )
        )
    return blocks, {"lines": len(blocks)}


PARSERS = {
    ".docx": _parse_docx,
    ".doc": _parse_docx,      # 老版 .doc 交给 docx 库尽力而为，失败会给明确提示
    ".pdf": _parse_pdf,
    ".txt": _parse_text,
    ".md": _parse_text,
}

SUPPORTED_EXT = sorted(PARSERS.keys())


# ---------------------------------------------------------------------------
# 对外入口
# ---------------------------------------------------------------------------

def parse_document(path: str | Path) -> ParsedDocument:
    """解析采购需求文件。"""
    p = Path(path)
    if not p.exists():
        raise ParseError(f"文件不存在：{p}")

    ext = p.suffix.lower()
    parser = PARSERS.get(ext)
    if parser is None:
        raise ParseError(
            f"不支持的文件格式：{ext}。当前支持：{'、'.join(SUPPORTED_EXT)}"
        )

    try:
        blocks, meta = parser(p)
    except ParseError:
        raise
    except Exception as e:
        hint = ""
        if ext == ".doc":
            hint = "（.doc 为老版格式，请用 Word 另存为 .docx 后重试）"
        raise ParseError(f"解析失败：{type(e).__name__}: {e}{hint}") from e

    if not blocks:
        raise ParseError("文件解析后内容为空，请检查文件是否有正文内容")

    doc = ParsedDocument(filename=p.name, filetype=ext, blocks=blocks, meta=meta)
    doc.meta["char_count"] = len(doc.full_text)
    doc.meta["section_list"] = list(doc.sections().keys())
    return doc


def chunk_text(text: str, size: int = 900, overlap: int = 120) -> list[str]:
    """把长文本切成带重叠的片段，用于知识库索引。"""
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []
    if len(text) <= size:
        return [text]

    chunks: list[str] = []
    start = 0
    n = len(text)
    while start < n:
        end = min(start + size, n)
        # 尽量在句子边界断开
        if end < n:
            window = text[start:end]
            for sep in ("。", "\n", "；", "，"):
                pos = window.rfind(sep)
                if pos > size * 0.5:
                    end = start + pos + 1
                    break
        piece = text[start:end].strip()
        if piece:
            chunks.append(piece)
        if end >= n:
            break
        start = max(end - overlap, start + 1)
    return chunks
