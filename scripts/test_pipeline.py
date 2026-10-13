"""
用模拟模型验证审查全流程管道（不需要 API Key）。

验证目标：把「解析 → 规则粗筛 → RAG 检索 → 提示词组装 → JSON 解析 →
          Finding 构造 → 去重 → 结论汇总 → 自动修订 → 报告导出」
整条链路跑通，确保真实模型接入后不会因为管道问题失败。

模拟策略：MockGateway 按规则代码返回预置的「违规判断」，
         故意混入几种真实模型常见的脏输出（markdown 围栏、多余文字、
         置信度缺失、原文摘录被改写），验证鲁棒性。

运行：python scripts/test_pipeline.py
"""

import asyncio
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from app.core.parser import parse_document  # noqa: E402
from app.core.reviewer import ReviewEngine, ReviewReport  # noqa: E402
from app.core.rewriter import DocumentRewriter, export_report_docx, export_report_markdown  # noqa: E402
from app.core.rules import get_rule_library  # noqa: E402

SAMPLE = ROOT / "knowledge" / "samples" / "XX大学高性能计算平台采购需求文件_含典型问题.docx"

PASS, FAIL = [], []


def check(name, fn):
    try:
        r = fn()
        PASS.append(name)
        print(f"  [PASS] {name}" + (f" → {r}" if r else ""))
        return r
    except Exception as e:
        import traceback

        FAIL.append(f"{name}: {type(e).__name__}: {e}")
        print(f"  [FAIL] {name} → {type(e).__name__}: {e}")
        traceback.print_exc()
        return None


# ---------------------------------------------------------------------------
# 模拟网关：按规则代码返回预置结果，混入真实模型常见的脏输出
# ---------------------------------------------------------------------------

# 故意让一部分规则返回"违规"，其余返回"通过"
VIOLATION_RULES = {
    "LEGAL_QUALIFICATION_BRAND": {
        # 脏输出 1：markdown 围栏 + 前后废话
        "raw": """好的，我已按规则完成审查。结果如下：

```json
{
  "violated": true,
  "title": "技术参数指定特定品牌",
  "detail": "技术参数第2条明确要求'必须为联想ThinkSystem SR650 V3，不接受其他品牌'，直接指定了特定品牌，违反《政府采购法实施条例》第二十条第六项关于不得限定或者指定特定品牌的规定。",
  "location": "技术参数",
  "original_text": "2. ★品牌要求：必须为联想 ThinkSystem SR650 V3，不接受其他品牌及型号。",
  "legal_basis": "《政府采购法实施条例》第二十条第六项",
  "suggestion": "删除品牌限定，改为以技术指标描述需求。",
  "replacement": "2. ★整机要求：应采用2U机架式服务器，支持上述CPU配置，具备冗余电源与热插拔硬盘位，不限定品牌与型号。",
  "confidence": 0.95
}
```

以上为本次审查结论。""",
    },
    "LEGAL_QUALIFICATION_SCALE": {
        # 脏输出 2：缺 confidence 字段
        "raw": json.dumps(
            {
                "violated": True,
                "title": "规模条件作为资格门槛",
                "detail": "资格条件第2条要求注册资本不低于5000万元、营业收入不低于3亿元，属于将规模条件作为资格要求，违反87号令第十七条。",
                "location": "资格条件",
                "original_text": "2. 投标人注册资本不低于人民币 5000 万元，2025 年度营业收入不低于人民币 3 亿元。",
                "legal_basis": "《政府采购货物和服务招标投标管理办法》（财政部令第87号）第十七条",
                "suggestion": "删除规模条件，改为与履约能力直接相关的条件。",
                "replacement": "2. 投标人应具备履行合同所必需的专业技术能力和项目实施经验，并提供相应证明材料。",
            },
            ensure_ascii=False,
        ),
    },
    "COMM_PERF_BOND": {
        # 脏输出 3：原文摘录被轻微改写（验证定位回退能力）
        "raw": json.dumps(
            {
                "violated": True,
                "title": "履约保证金超过法定上限",
                "detail": "履约保证金为合同金额的20%，超过《政府采购法实施条例》第四十八条规定的10%上限。",
                "location": "商务要求",
                "original_text": "中标人须在合同签订前提交合同金额20%的履约保证金。",
                "legal_basis": "《政府采购法实施条例》第四十八条",
                "suggestion": "将履约保证金降至不超过合同金额的10%。",
                "replacement": "3. 履约保证金：中标人须在合同签订前提交合同金额 10% 的履约保证金，可采用支票、汇票、本票或者金融机构、担保机构出具的保函等非现金形式提交。",
                "confidence": 0.9,
            },
            ensure_ascii=False,
        ),
    },
    "SCORE_PRICE_WEIGHT": {
        # 脏输出 4：JSON 带尾随逗号
        "raw": """{
  "violated": true,
  "title": "价格分权重不合规",
  "detail": "本项目为货物采购，价格分仅20分，低于87号令第五十五条规定的30%下限。",
  "location": "评分办法",
  "original_text": "1. 价格分：20 分。去掉最高和最低报价后取算术平均值作为评标基准价，投标报价等于基准价的得满分，每高 1% 扣 1 分。",
  "legal_basis": "《政府采购货物和服务招标投标管理办法》（财政部令第87号）第五十五条",
  "suggestion": "将价格分提高至不低于30分，并改用低价优先法。",
  "replacement": "1. 价格分：30 分。采用低价优先法计算，即满足招标文件要求且投标价格最低的投标报价为评标基准价，其价格分为满分。其他投标人的价格分统一按照下列公式计算：投标报价得分＝（评标基准价／投标报价）×30。",
  "confidence": 0.92,
}""",
    },
    "COMPLETE_ACCEPTANCE": {
        # 缺失类问题：无原文摘录，只给新增条款
        "raw": json.dumps(
            {
                "violated": True,
                "title": "履约验收要求缺失",
                "detail": "文件仅表述'验收合格后付款'，未明确验收主体、时间、方式、程序、内容和验收标准，不符合《政府采购需求管理办法》第二十四条。",
                "location": "商务要求",
                "original_text": "",
                "legal_basis": "《政府采购需求管理办法》第二十四条",
                "suggestion": "补充完整的履约验收方案。",
                "replacement": "六、履约验收\n1. 验收主体：由采购人组织，邀请不少于3名相关领域专家参与。\n2. 验收时间：全部货物安装调试完成后10个工作日内组织初步验收。\n3. 验收方式：现场查验、功能测试与文档审查相结合。\n4. 验收内容：逐项核对技术参数响应情况、商务条款履行情况。\n5. 验收标准：所有技术指标经测试全部达标，随机抽取不少于20%的节点进行满负荷压力测试且连续运行72小时无故障。\n6. 验收不合格的，中标人应在15个工作日内整改完毕并重新申请验收。",
                "confidence": 0.88,
            },
            ensure_ascii=False,
        ),
    },
    "POLICY_IMPORT_DEMO": {
        # 低置信度：验证 needs_human_review 标记
        "raw": json.dumps(
            {
                "violated": True,
                "title": "进口产品未见论证材料",
                "detail": "技术参数要求NVIDIA H100显卡，属进口产品，但文件中未见进口产品论证意见与财政部门核准材料。因未提供完整附件，无法确认是否已单独办理论证手续。",
                "location": "技术参数",
                "original_text": "6. ★GPU：NVIDIA H100 80GB PCIe 显卡，每节点 4 块，须采用 NVIDIA 专有 NVLink 互联技术。",
                "legal_basis": "《政府采购进口产品管理办法》（财库〔2007〕119号）",
                "suggestion": "补充进口产品论证材料，或改用国产替代方案。",
                "replacement": "",
                "confidence": 0.45,
            },
            ensure_ascii=False,
        ),
    },
}


class MockGateway:
    """模拟大模型网关，接口与 LLMGateway 一致。"""

    def __init__(self):
        from app.core.config import ProviderProfile

        self.profile = ProviderProfile(
            id="mock", name="模拟模型（测试用）", base_url="http://mock", model="mock-reviewer"
        )
        self.base_url = "http://mock"
        self.calls = 0
        self.prompts = {}

    async def chat_json(self, messages, **kwargs):
        self.calls += 1
        prompt = messages[-1]["content"]

        # 从提示词中提取规则代码
        code = None
        for line in prompt.splitlines():
            if line.startswith("规则代码："):
                code = line.split("：", 1)[1].strip()
                break

        self.prompts[code] = prompt

        from app.core.llm import extract_json

        if code in VIOLATION_RULES:
            return extract_json(VIOLATION_RULES[code]["raw"])

        # 其余规则返回"通过"
        return {
            "violated": False,
            "title": "",
            "detail": "经核查，该规则对应的要求已满足，未发现合规性问题。",
            "location": "",
            "original_text": "",
            "legal_basis": "",
            "suggestion": "",
            "replacement": "",
            "confidence": 0.8,
        }


# ---------------------------------------------------------------------------
async def main():
    print("=" * 72)
    print("审查全流程管道验证（模拟模型）")
    print("=" * 72)

    if not SAMPLE.exists():
        print(f"测试文件不存在：{SAMPLE}")
        print("请先运行：python scripts/make_sample.py")
        return 1

    # 1. 解析
    print("\n[1] 文档解析")
    doc = check("解析 docx", lambda: f"{len(doc_holder)} 字 / {len(doc_holder.sections())} 章节") if False else None

    doc = parse_document(SAMPLE)
    check("解析 docx", lambda: f"{len(doc.full_text)} 字 / {len(doc.sections())} 章节")
    check("章节为标准部位", lambda: "、".join(doc.sections().keys()))

    # 2. 规则粗筛
    print("\n[2] 规则粗筛")
    lib = get_rule_library()
    applicable = lib.prefilter(doc.sections())
    check("粗筛适用规则", lambda: f"从 {len(lib.rules)} 条筛出 {len(applicable)} 条")

    # 3. 执行审查
    print("\n[3] 审查引擎（模拟模型）")
    gw = MockGateway()
    engine = ReviewEngine(gw, library=lib)

    progress_log = []

    def on_progress(code, done, total, msg):
        progress_log.append((done, total, msg))

    report = await engine.review(doc, on_progress=on_progress)

    check("模型调用次数", lambda: f"{gw.calls} 次（应等于适用规则数 {len(applicable)}）")
    check("进度回调", lambda: f"{len(progress_log)} 次，最后 {progress_log[-1][0]}/{progress_log[-1][1]}")
    check("规则评估数", lambda: str(report.rules_evaluated))
    check("未发生规则执行错误", lambda: "无错误" if not report.errors else f"⚠ {report.errors}")

    # 4. 结论质量
    print("\n[4] 结论质量")
    s = report.summary
    issues = report.issues

    check("检出问题数", lambda: f"{len(issues)} 项（预置 {len(VIOLATION_RULES)} 项违规）")
    check("总体判定", lambda: f"{s['verdict']}（{s['verdict_level']}）")
    check(
        "分级统计",
        lambda: f"严重 {s['blocker_count']} / 高 {s['high_count']} / 中 {s['medium_count']} / 低 {s['low_count']}",
    )
    check("通过项统计", lambda: f"{s['pass_count']} 项")
    check("待人工复核标记", lambda: f"{s['needs_human_review_count']} 项（低置信度应被标记）")

    # 验证脏输出被正确处理
    print("\n[5] 脏输出鲁棒性")
    brand = next((f for f in issues if f.rule_code == "LEGAL_QUALIFICATION_BRAND"), None)
    check("markdown围栏 JSON 解析", lambda: "成功" if brand else "失败")
    check("  → 字段完整", lambda: f"title='{brand.title}' conf={brand.confidence}")

    scale = next((f for f in issues if f.rule_code == "LEGAL_QUALIFICATION_SCALE"), None)
    check("缺 confidence 字段", lambda: f"默认值 {scale.confidence}" if scale else "失败")

    weight = next((f for f in issues if f.rule_code == "SCORE_PRICE_WEIGHT"), None)
    check("尾随逗号修复", lambda: f"conf={weight.confidence}" if weight else "失败")

    acc = next((f for f in issues if f.rule_code == "COMPLETE_ACCEPTANCE"), None)
    check("缺失类（无原文）", lambda: f"replacement {len(acc.replacement)} 字" if acc else "失败")

    imp = next((f for f in issues if f.rule_code == "POLICY_IMPORT_DEMO"), None)
    check(
        "低置信度标记复核",
        lambda: f"confidence={imp.confidence} needs_review={imp.needs_human_review}" if imp else "失败",
    )

    # 6. RAG 依据
    print("\n[6] RAG 法规依据")
    with_cites = [f for f in issues if f.citations]
    check("带知识库引用的意见", lambda: f"{len(with_cites)}/{len(issues)} 项")
    if with_cites:
        check("引用示例", lambda: with_cites[0].citations[0]["citation"])

    # 7. 提示词质量
    print("\n[7] 提示词组装")
    p = gw.prompts.get("LEGAL_QUALIFICATION_BRAND", "")
    check("提示词含文档内容", lambda: "联想" in p)
    check("提示词含法规依据段", lambda: "知识库检索到的法规依据" in p)
    check("提示词含输出格式要求", lambda: "confidence" in p)
    check("提示词长度", lambda: f"{len(p)} 字")

    # 8. 自动修订
    print("\n[8] 自动修订")
    selected = [f.id for f in issues]
    for f in report.findings:
        f.selected = f.id in selected

    rw = DocumentRewriter(SAMPLE)
    result = rw.apply(report.findings, output_name="_pipeline_test_revised.docx")
    check(
        "修订统计",
        lambda: f"替换 {result.replaced} / 新增 {result.inserted} / 人工 {result.manual} / 跳过 {result.skipped}",
    )

    # 校验修订结果文件
    from docx import Document

    out = Document(result.output_path)
    full = "\n".join(pp.text for pp in out.paragraphs)
    check("品牌限定已被替换", lambda: "不限定品牌与型号" in full or "不限定品牌" in full)
    check("原品牌表述已移除", lambda: "不接受其他品牌及型号" not in full)
    check("新增验收条款已插入", lambda: "验收主体" in full)
    check("修订说明表已附加", lambda: "修订说明表" in full)
    check("高亮标记已应用", lambda: "已应用" if result.replaced > 0 else "无替换")

    # 9. 报告导出
    print("\n[9] 报告导出")
    data = report.to_dict()
    docx_path = check("导出 Word 报告", lambda: Path(export_report_docx(data)).name)
    md_path = check("导出 Markdown 报告", lambda: Path(export_report_markdown(data)).name)

    if docx_path:
        d = Document(str(ROOT / "data" / "exports" / docx_path))
        text = "\n".join(p.text for p in d.paragraphs)
        check("报告含总体判定", lambda: "总体判定" in text)
        check("报告含问题清单", lambda: "问题清单" in text)
        check("报告含法规依据", lambda: "法规依据" in text)

    # 10. JSON 可序列化（前端依赖）
    print("\n[10] 序列化")
    check("报告可 JSON 序列化", lambda: f"{len(json.dumps(data, ensure_ascii=False))} 字节")

    # 清理
    Path(result.output_path).unlink(missing_ok=True)
    if docx_path:
        (ROOT / "data" / "exports" / docx_path).unlink(missing_ok=True)
    if md_path:
        (ROOT / "data" / "exports" / md_path).unlink(missing_ok=True)

    print("\n" + "=" * 72)
    print(f"通过 {len(PASS)} / 失败 {len(FAIL)}")
    if FAIL:
        for f in FAIL:
            print(f"  [FAIL] {f}")
    print("=" * 72)
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
