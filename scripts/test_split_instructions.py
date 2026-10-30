"""回归测试：inspect 必须能解析被拆分成多个 run 的字段指令。

Word 打开并保存文档后，会把一条长 instrText 拆进多个 run。
修复前的正则要求 JSON 与 </w:instrText> 相邻，于是全部字段解析失败
（表现为 json_failures == item_fields）。

本测试自造样本，不依赖任何外部文档：先用 zotero_fields 生成一份带引文字段与
参考文献字段的 docx，再把每条 instrText 切成多段，验证解析结果不变。
"""

from __future__ import annotations

import json
import re
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import zotero_fields as Z  # noqa: E402
from docx import Document  # noqa: E402

STYLE = "http://www.zotero.org/styles/china-national-standard-gb-t-7714-2015-numeric"
INSTR = re.compile(r"(<w:instrText[^>]*>)(.*?)(</w:instrText>)", re.S)

ITEMS = [
    {
        "key": "SYNTH001",
        "itemType": "journalArticle",
        "title": "合成测试文献一",
        "creators": ["张三", "李四"],
        "publicationTitle": "测试期刊",
        "date": "2026",
        "volume": "12",
        "issue": "3",
        "pages": "1-9",
        "DOI": "",
        "url": "",
    },
    {
        "key": "SYNTH002",
        "itemType": "journalArticle",
        "title": "合成测试文献二",
        "creators": ["王五"],
        "publicationTitle": "测试期刊",
        "date": "2025",
        "volume": "11",
        "issue": "2",
        "pages": "10-18",
        "DOI": "",
        "url": "",
    },
]


def build_sample(path: Path) -> None:
    """生成一份含 2 个引文字段与 1 个参考文献字段的 docx。"""
    doc = Document()
    para = doc.add_paragraph("测试段落：")
    for index, item in enumerate(ITEMS, start=1):
        csl = Z.build_csl_item(item, 1000 + index, f"http://zotero.org/users/1/items/{item['key']}")
        citation = Z.build_citation_json(csl, f"[{index}]", f"cit{index}")
        Z.add_citation_field(para, citation, f"[{index}]")

    doc.add_paragraph("参考文献")
    refs = []
    for index, item in enumerate(ITEMS, start=1):
        refs.append(doc.add_paragraph(f"[{index}] {item['title']}"))
    Z.add_bibl_field(refs)

    doc.save(str(path))
    Z.inject_zotero_prefs(path, STYLE)


def split_instructions(xml: str, parts: int = 3) -> str:
    """把每条 instrText 的内容切成长度相近的若干段，模拟 Word 的拆分。"""

    def repl(match: re.Match) -> str:
        open_tag, body, close_tag = match.group(1), match.group(2), match.group(3)
        if len(body) < parts * 2:
            return match.group(0)
        size = len(body) // parts
        pieces = [body[i : i + size] for i in range(0, len(body), size)]
        return "".join(f"{open_tag}{piece}{close_tag}" for piece in pieces if piece)

    return INSTR.sub(repl, xml)


def main() -> int:
    work = Path(tempfile.mkdtemp(prefix="zot_split_"))
    sample = work / "sample.docx"
    build_sample(sample)

    base = Z.inspect(sample)
    print("原样本：")
    print(
        f"  item_fields={base['item_fields']} parsed={base['parsed_ok']} "
        f"failures={base['json_failures']} unique={base['unique_cited_items']}"
    )

    with zipfile.ZipFile(sample) as package:
        entries = {name: package.read(name) for name in package.namelist()}

    xml = entries["word/document.xml"].decode("utf-8")
    before = len(INSTR.findall(xml))
    xml = split_instructions(xml)
    after = len(INSTR.findall(xml))
    entries["word/document.xml"] = xml.encode("utf-8")
    print(f"\n拆分指令 run：{before} -> {after}")

    split_path = work / "split.docx"
    with zipfile.ZipFile(split_path, "w", zipfile.ZIP_DEFLATED) as package:
        for name, data in entries.items():
            package.writestr(name, data)

    split = Z.inspect(split_path)
    print("拆分后：")
    print(
        f"  item_fields={split['item_fields']} parsed={split['parsed_ok']} "
        f"failures={split['json_failures']} unique={split['unique_cited_items']}"
    )

    checks = [
        ("样本含 2 个引文字段", base["item_fields"] == 2),
        ("样本 JSON 全部可解析", base["json_failures"] == 0),
        ("拆分后仍识别出全部字段", split["item_fields"] == base["item_fields"]),
        ("拆分后 JSON 全部可解析", split["json_failures"] == 0),
        ("拆分后唯一条目数一致", split["unique_cited_items"] == base["unique_cited_items"]),
        ("参考文献字段仍在", split["bibl_fields"] == 1),
        ("字段标记数一致", split["fldChar_begin"] == base["fldChar_begin"]),
    ]
    print()
    failed = 0
    for label, ok in checks:
        print(f"  [{'PASS' if ok else 'FAIL'}] {label}")
        failed += 0 if ok else 1

    print(f"\n临时文件：{split_path}")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
