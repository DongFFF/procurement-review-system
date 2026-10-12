"""生成测试用采购需求文件（含典型违规点），用于验证审查与修订全流程。

运行：python scripts/make_sample.py
产物：knowledge/samples/XX大学高性能计算平台采购需求文件_含典型问题.docx
"""

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "knowledge" / "samples"
OUT_DIR.mkdir(parents=True, exist_ok=True)

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
except ImportError:
    print("需要 python-docx：pip install python-docx")
    sys.exit(1)


def build() -> Path:
    doc = Document()

    # 标题
    t = doc.add_paragraph()
    tr = t.add_run("XX大学高性能计算平台采购需求文件")
    tr.bold = True
    tr.font.size = Pt(18)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sr = sub.add_run("（项目编号：XDCG-2026-001）")
    sr.font.size = Pt(10.5)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 一、项目概况
    doc.add_heading("一、项目概况", level=1)
    doc.add_paragraph(
        "为满足我校人工智能与高性能计算科研平台建设需要，拟采购高性能计算集群一套，"
        "用于深度学习模型训练与科学计算。本项目预算金额为人民币 860 万元，"
        "资金来源为中央高校基本科研业务费及一流学科建设经费。"
    )
    doc.add_paragraph(
        "采购内容包括：计算节点 32 台、管理节点 2 台、存储系统 1 套、"
        "InfiniBand 交换机 4 台、机柜及配套线缆若干。"
    )

    # 二、技术参数
    doc.add_heading("二、技术参数要求", level=1)
    doc.add_paragraph("（一）计算节点（32 台）")
    doc.add_paragraph("1. ★CPU：Intel Xeon Platinum 8480+ 处理器，2 颗，主频不低于 2.0GHz，"
                      "核心数不少于 56 核。")
    doc.add_paragraph("2. ★品牌要求：必须为联想 ThinkSystem SR650 V3，不接受其他品牌及型号。")
    doc.add_paragraph("3. 内存：512GB DDR5 ECC RDIMM，频率不低于 4800MHz。")
    doc.add_paragraph("4. 系统盘：2 块 960GB NVMe SSD 组成 RAID1。")
    doc.add_paragraph("5. 数据盘：8 块 3.84TB NVMe SSD。")
    doc.add_paragraph("6. ★GPU：NVIDIA H100 80GB PCIe 显卡，每节点 4 块，"
                      "须采用 NVIDIA 专有 NVLink 互联技术。")
    doc.add_paragraph("7. 网络：配置 2 端口 100Gb InfiniBand HDR 网卡。")
    doc.add_paragraph("8. 电源：冗余电源，功率不超过 2000W。")
    doc.add_paragraph("9. 机箱尺寸须为（高×宽×深）87mm×447mm×780mm，重量不超过 35kg。")

    doc.add_paragraph("（二）存储系统")
    doc.add_paragraph("1. 裸容量不低于 1.5PB，采用分布式并行文件系统。")
    doc.add_paragraph("2. 须支持基于 XX 公司专利的智能分层存储算法，"
                      "并具备与我校现有 XX 品牌存储平台无缝对接的私有协议接口。")
    doc.add_paragraph("3. 缓存容量不低于 512GB。")

    # 三、资格条件
    doc.add_heading("三、供应商资格条件", level=1)
    doc.add_paragraph("1. 投标人须为在中华人民共和国境内注册的独立法人，"
                      "且为国有独资或国有控股企业。")
    doc.add_paragraph("2. 投标人注册资本不低于人民币 5000 万元，"
                      "2025 年度营业收入不低于人民币 3 亿元。")
    doc.add_paragraph("3. 投标人须提供近三年在北京市高校的同类项目业绩不少于 3 份，"
                      "每提供 1 份加 1 分。")
    doc.add_paragraph("4. 投标人须提供所投产品生产厂家针对本项目出具的唯一专项授权书原件，"
                      "否则投标无效。")
    doc.add_paragraph("5. 投标人须为中国教育装备行业协会会员单位，"
                      "且在北京市拥有不少于 500 平方米的办公场所。")
    doc.add_paragraph("6. 本项目仅接受已入围我校“信息化设备供应商预选名录”的单位投标。")
    doc.add_paragraph("7. 投标人及其法定代表人近五年内有任何行政处罚记录的，不得参与本项目。")

    # 四、商务要求
    doc.add_heading("四、商务要求", level=1)
    doc.add_paragraph("1. 交货期：合同签订后按采购人要求的时间交货；"
                      "交货地点由采购人指定。")
    doc.add_paragraph("2. 投标保证金：人民币 30 万元。")
    doc.add_paragraph("3. 履约保证金：中标人须在合同签订前提交合同金额 20% 的履约保证金，"
                      "且须以现金或电汇方式提交。")
    doc.add_paragraph("4. 付款方式：项目全部完成并验收合格，"
                      "经审计部门审计完成后，于次年度预算下达后一次性支付合同价款的 100%，"
                      "不设预付款。")
    doc.add_paragraph("5. 质保期：投标人须承诺对所投全部设备提供不少于 8 年免费质保"
                      "及终身免费上门维护，质保期内免费更换一切零部件（含耗材）。")
    doc.add_paragraph("6. 违约责任：供应商逾期交付的，每逾期一日按合同总额的 5% 支付违约金，"
                      "上不封顶；采购人逾期付款的，不承担违约责任。")
    doc.add_paragraph("7. 投标报价低于项目预算金额 85% 的，视为恶意低价，作无效投标处理。")
    doc.add_paragraph("8. 本项目不接受联合体投标。")

    # 五、评分办法
    doc.add_heading("五、评分办法", level=1)
    doc.add_paragraph("本项目采用综合评分法，总分 100 分，具体如下：")
    doc.add_paragraph("1. 价格分：20 分。去掉最高和最低报价后取算术平均值作为评标基准价，"
                      "投标报价等于基准价的得满分，每高 1% 扣 1 分。")
    doc.add_paragraph("2. 技术方案：25 分。根据投标人技术方案的先进性、合理性综合打分，"
                      "优得 20—25 分，良得 15—19 分，中得 8—14 分，差得 0—7 分。")
    doc.add_paragraph("3. 企业综合实力：15 分。根据投标人企业规模、行业影响力综合评定。")
    doc.add_paragraph("4. 服务方案：20 分。根据投标人售后服务方案的完善程度综合评定。")
    doc.add_paragraph("5. 类似业绩：10 分。投标人近三年获得北京市“优秀供应商”称号的，"
                      "每项加 2 分，最多加 6 分；具有信息系统集成一级资质的加 4 分。")
    doc.add_paragraph("6. 节能环保：10 分。投标产品具有节能产品认证证书的得 5 分，"
                      "具有环境标志产品认证证书的得 5 分。")
    doc.add_paragraph("注：具体评分细则由评标委员会在评标现场根据项目情况确定。")

    # 六、履约验收
    doc.add_heading("六、履约验收", level=1)
    doc.add_paragraph("项目完成后由采购人组织验收，验收合格后付款。")

    out = OUT_DIR / "XX大学高性能计算平台采购需求文件_含典型问题.docx"
    doc.save(str(out))
    return out


if __name__ == "__main__":
    p = build()
    print(f"已生成测试文件：{p}")
    print(f"大小：{p.stat().st_size} 字节")
    print()
    print("该文件刻意植入了以下典型违规点，用于验证系统检出能力：")
    print("  - 指定品牌型号（联想 ThinkSystem SR650 V3）")
    print("  - 指向特定专利技术路线（NVIDIA NVLink、XX 公司专利算法）")
    print("  - 限定所有制形式（国有独资或国有控股）")
    print("  - 规模条件设门槛（注册资本 5000 万、营业收入 3 亿）")
    print("  - 特定区域业绩限制（北京市高校业绩）")
    print("  - 厂家授权作资格门槛")
    print("  - 与履约无关的资格条件（行业协会会员、办公场所面积）")
    print("  - 违规设置供应商库（预选名录）")
    print("  - 投标保证金超额（30 万 / 860 万 = 3.5% > 2%）")
    print("  - 履约保证金超额（20% > 10%）")
    print("  - 付款期限过长、要求全额垫资")
    print("  - 质保期过长（8 年 + 终身免费）")
    print("  - 违约责任不对等")
    print("  - 设定最低限价（低于预算 85% 作废标）")
    print("  - 价格分权重不合规（货物 20% < 30%）")
    print("  - 价格分算法违规（去掉最高最低报价）")
    print("  - 主观分占比过高（60 分）")
    print("  - 区域行业奖项加分")
    print("  - 资格条件充作评分项（信息系统集成资质）")
    print("  - 评标标准未事先载明（评标现场确定细则）")
    print("  - 履约验收要求缺失")
    print("  - 进口产品未论证（NVIDIA H100）")
